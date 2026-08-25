from __future__ import annotations

import os
from dataclasses import dataclass, field
from multiprocessing import get_context
from multiprocessing.connection import Connection
from multiprocessing.process import BaseProcess
from multiprocessing.synchronize import Event as EventType
from pathlib import Path
from time import monotonic
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from pypdf import PdfReader


class ParserError(RuntimeError):
    pass


@dataclass(frozen=True)
class ParsedBlock:
    text: str
    locator: dict[str, object]
    metadata: dict[str, object]


@dataclass(frozen=True)
class ParsedDocument:
    media_type: str
    blocks: tuple[ParsedBlock, ...]


@dataclass(frozen=True)
class _ParserLimits:
    maximum_document_characters: int
    maximum_blocks: int
    maximum_pdf_pages: int
    maximum_docx_expanded_bytes: int
    maximum_parse_seconds: float
    maximum_worker_memory_bytes: int


@dataclass
class _BlockCollector:
    maximum_characters: int
    maximum_blocks: int
    blocks: list[ParsedBlock] = field(default_factory=list)
    character_count: int = 0

    def append(self, block: ParsedBlock) -> None:
        if len(self.blocks) >= self.maximum_blocks:
            raise ParserError("parser_resource_limit")
        self.character_count += len(block.text)
        if self.character_count > self.maximum_characters:
            raise ParserError("parser_resource_limit")
        self.blocks.append(block)


def parse_document(
    path: Path,
    media_type: str,
    *,
    maximum_document_characters: int = 5_000_000,
    maximum_blocks: int = 20_000,
    maximum_pdf_pages: int = 2_000,
    maximum_docx_expanded_bytes: int = 100_000_000,
    maximum_parse_seconds: float = 30.0,
    maximum_worker_memory_bytes: int = 512_000_000,
) -> ParsedDocument:
    if min(maximum_document_characters, maximum_blocks, maximum_pdf_pages) < 1:
        raise ValueError("parser limits must be positive")
    if min(maximum_docx_expanded_bytes, maximum_worker_memory_bytes) < 1:
        raise ValueError("parser limits must be positive")
    if maximum_parse_seconds <= 0:
        raise ValueError("parser limits must be positive")
    limits = _ParserLimits(
        maximum_document_characters=maximum_document_characters,
        maximum_blocks=maximum_blocks,
        maximum_pdf_pages=maximum_pdf_pages,
        maximum_docx_expanded_bytes=maximum_docx_expanded_bytes,
        maximum_parse_seconds=maximum_parse_seconds,
        maximum_worker_memory_bytes=maximum_worker_memory_bytes,
    )
    context = get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    start_event = context.Event()
    process = context.Process(
        target=_parser_process,
        args=(sender, start_event, path, media_type, limits),
        daemon=True,
    )
    process.start()
    sender.close()
    job_handle: int | None = None
    try:
        if os.name == "nt":
            job_handle = _assign_windows_memory_job(
                process, maximum_worker_memory_bytes
            )
        start_event.set()
        if not receiver.poll(maximum_parse_seconds):
            _terminate_parser(process)
            raise ParserError("parser_resource_limit")
        status, payload = receiver.recv()
        if status != "ok":
            raise ParserError(str(payload))
        if not isinstance(payload, ParsedDocument):
            raise ParserError("parser_failed")
        return payload
    except (EOFError, OSError) as error:
        _terminate_parser(process)
        raise ParserError("parser_failed") from error
    finally:
        receiver.close()
        process.join(timeout=1)
        if process.is_alive():
            _terminate_parser(process)
        if job_handle is not None:
            _close_windows_handle(job_handle)


def _parser_process(
    sender: Connection,
    start_event: EventType,
    path: Path,
    media_type: str,
    limits: _ParserLimits,
) -> None:
    try:
        if not start_event.wait(timeout=10):
            raise ParserError("parser_resource_limit")
        _apply_memory_limit(limits.maximum_worker_memory_bytes)
        sender.send(("ok", _parse_document_in_process(path, media_type, limits)))
    except ParserError as error:
        sender.send(("error", str(error)))
    except BaseException:
        sender.send(("error", "parser_failed"))
    finally:
        sender.close()


def _parse_document_in_process(
    path: Path, media_type: str, limits: _ParserLimits
) -> ParsedDocument:
    deadline = monotonic() + limits.maximum_parse_seconds
    collector = _BlockCollector(
        maximum_characters=limits.maximum_document_characters,
        maximum_blocks=limits.maximum_blocks,
    )
    try:
        if media_type == "application/pdf":
            _parse_pdf(path, limits.maximum_pdf_pages, deadline, collector)
        elif media_type == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            _parse_docx(
                path,
                limits.maximum_docx_expanded_bytes,
                deadline,
                collector,
            )
        elif media_type == "text/markdown":
            _parse_markdown(path, deadline, collector)
        elif media_type == "text/plain":
            _parse_plain_text(path, deadline, collector)
        else:
            raise ParserError("unsupported_media_type")
    except ParserError:
        raise
    except Exception as error:
        raise ParserError("parser_failed") from error
    if not collector.blocks:
        raise ParserError("parser_empty")
    return ParsedDocument(media_type=media_type, blocks=tuple(collector.blocks))


def _parse_pdf(
    path: Path,
    maximum_pages: int,
    deadline: float,
    collector: _BlockCollector,
) -> None:
    pages = PdfReader(path).pages
    if len(pages) > maximum_pages:
        raise ParserError("parser_resource_limit")
    for page_number, page in enumerate(pages, start=1):
        _check_deadline(deadline)
        text = (page.extract_text() or "").strip()
        if text:
            collector.append(
                ParsedBlock(
                    text=text,
                    locator={"page": page_number},
                    metadata={"kind": "page"},
                )
            )


def _parse_docx(
    path: Path,
    maximum_expanded_bytes: int,
    deadline: float,
    collector: _BlockCollector,
) -> None:
    try:
        with ZipFile(path) as archive:
            expanded_bytes = sum(info.file_size for info in archive.infolist())
    except BadZipFile as error:
        raise ParserError("parser_failed") from error
    if expanded_bytes > maximum_expanded_bytes:
        raise ParserError("parser_resource_limit")
    heading: str | None = None
    paragraph_number = 0
    table_number = 0
    for item in Document(path).iter_inner_content():
        _check_deadline(deadline)
        if isinstance(item, Paragraph):
            paragraph_number += 1
            block, heading = _docx_paragraph_block(
                item,
                locator={"paragraph": paragraph_number},
                heading=heading,
            )
            if block is not None:
                collector.append(block)
            continue
        if not isinstance(item, Table):
            continue
        table_number += 1
        for row_number, row in enumerate(item.rows, start=1):
            for cell_number, cell in enumerate(row.cells, start=1):
                for cell_paragraph_number, paragraph in enumerate(
                    cell.paragraphs, start=1
                ):
                    _check_deadline(deadline)
                    block, heading = _docx_paragraph_block(
                        paragraph,
                        locator={
                            "table": table_number,
                            "row": row_number,
                            "cell": cell_number,
                            "paragraph": cell_paragraph_number,
                        },
                        heading=heading,
                    )
                    if block is not None:
                        collector.append(block)


def _docx_paragraph_block(
    paragraph: Paragraph, *, locator: dict[str, object], heading: str | None
) -> tuple[ParsedBlock | None, str | None]:
    text = paragraph.text.strip()
    if not text:
        return None, heading
    style_name = paragraph.style.name if paragraph.style else ""
    is_heading = style_name.startswith("Heading")
    if is_heading:
        heading = text
    if heading:
        locator["heading"] = heading
    return (
        ParsedBlock(
            text=text,
            locator=locator,
            metadata={"kind": "heading" if is_heading else "paragraph"},
        ),
        heading,
    )


def _parse_markdown(
    path: Path, deadline: float, collector: _BlockCollector
) -> None:
    text = path.read_text(encoding="utf-8")
    tokens = MarkdownIt("commonmark").parse(text)
    heading: str | None = None
    next_inline_is_heading = False
    for token in tokens:
        _check_deadline(deadline)
        if token.type == "heading_open":
            next_inline_is_heading = True
            continue
        if token.type not in {"inline", "fence", "code_block"}:
            continue
        is_code = token.type in {"fence", "code_block"}
        content = token.content.rstrip("\r\n") if is_code else token.content.strip()
        if not content:
            continue
        if next_inline_is_heading:
            heading = content
            kind = "heading"
            next_inline_is_heading = False
        else:
            kind = "code" if is_code else "paragraph"
        locator: dict[str, object] = {}
        if token.map:
            locator.update(
                {"line_start": token.map[0] + 1, "line_end": token.map[1]}
            )
        if heading:
            locator["heading"] = heading
        collector.append(
            ParsedBlock(text=content, locator=locator, metadata={"kind": kind})
        )


def _parse_plain_text(
    path: Path, deadline: float, collector: _BlockCollector
) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    start: int | None = None
    collected: list[str] = []
    for line_number, line in enumerate([*lines, ""], start=1):
        _check_deadline(deadline)
        if line.strip():
            if start is None:
                start = line_number
            collected.append(line)
            continue
        if collected and start is not None:
            collector.append(
                ParsedBlock(
                    text="\n".join(collected),
                    locator={"line_start": start, "line_end": line_number - 1},
                    metadata={"kind": "paragraph"},
                )
            )
        start = None
        collected = []


def _check_deadline(deadline: float) -> None:
    if monotonic() > deadline:
        raise ParserError("parser_resource_limit")


def _terminate_parser(process: BaseProcess) -> None:
    process.terminate()
    process.join(timeout=5)
    if process.is_alive():
        process.kill()
        process.join(timeout=5)


def _apply_memory_limit(maximum_bytes: int) -> None:
    if os.name == "nt":
        return
    try:
        import resource
    except ImportError:
        return
    try:
        resource.setrlimit(resource.RLIMIT_AS, (maximum_bytes, maximum_bytes))
    except (OSError, ValueError):
        raise ParserError("parser_resource_limit")


def _assign_windows_memory_job(process: BaseProcess, maximum_bytes: int) -> int:
    import ctypes
    from ctypes import wintypes

    class _IoCounters(ctypes.Structure):
        _fields_ = [
            ("ReadOperationCount", ctypes.c_ulonglong),
            ("WriteOperationCount", ctypes.c_ulonglong),
            ("OtherOperationCount", ctypes.c_ulonglong),
            ("ReadTransferCount", ctypes.c_ulonglong),
            ("WriteTransferCount", ctypes.c_ulonglong),
            ("OtherTransferCount", ctypes.c_ulonglong),
        ]

    class _BasicLimitInformation(ctypes.Structure):
        _fields_ = [
            ("PerProcessUserTimeLimit", ctypes.c_longlong),
            ("PerJobUserTimeLimit", ctypes.c_longlong),
            ("LimitFlags", wintypes.DWORD),
            ("MinimumWorkingSetSize", ctypes.c_size_t),
            ("MaximumWorkingSetSize", ctypes.c_size_t),
            ("ActiveProcessLimit", wintypes.DWORD),
            ("Affinity", ctypes.c_size_t),
            ("PriorityClass", wintypes.DWORD),
            ("SchedulingClass", wintypes.DWORD),
        ]

    class _ExtendedLimitInformation(ctypes.Structure):
        _fields_ = [
            ("BasicLimitInformation", _BasicLimitInformation),
            ("IoInfo", _IoCounters),
            ("ProcessMemoryLimit", ctypes.c_size_t),
            ("JobMemoryLimit", ctypes.c_size_t),
            ("PeakProcessMemoryUsed", ctypes.c_size_t),
            ("PeakJobMemoryUsed", ctypes.c_size_t),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateJobObjectW.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR]
    kernel32.CreateJobObjectW.restype = wintypes.HANDLE
    kernel32.SetInformationJobObject.argtypes = [
        wintypes.HANDLE,
        ctypes.c_int,
        ctypes.c_void_p,
        wintypes.DWORD,
    ]
    kernel32.SetInformationJobObject.restype = wintypes.BOOL
    kernel32.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel32.OpenProcess.restype = wintypes.HANDLE
    kernel32.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
    kernel32.AssignProcessToJobObject.restype = wintypes.BOOL
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL

    job = kernel32.CreateJobObjectW(None, None)
    if not job:
        _terminate_parser(process)
        raise ParserError("parser_resource_limit")
    information = _ExtendedLimitInformation()
    information.BasicLimitInformation.LimitFlags = 0x100 | 0x2000
    information.ProcessMemoryLimit = maximum_bytes
    configured = kernel32.SetInformationJobObject(
        job, 9, ctypes.byref(information), ctypes.sizeof(information)
    )
    process_handle = kernel32.OpenProcess(
        0x0001 | 0x0100 | 0x1000, False, int(process.pid or 0)
    )
    assigned = bool(
        process_handle and kernel32.AssignProcessToJobObject(job, process_handle)
    )
    if process_handle:
        kernel32.CloseHandle(process_handle)
    if not configured or not assigned:
        kernel32.CloseHandle(job)
        _terminate_parser(process)
        raise ParserError("parser_resource_limit")
    return int(job)


def _close_windows_handle(handle: int) -> None:
    import ctypes
    from ctypes import wintypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL
    kernel32.CloseHandle(wintypes.HANDLE(handle))
