from __future__ import annotations

import hashlib
import os
import time
from multiprocessing import active_children
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from spaces.learning.services.db.models import (
    Base,
    LearningInvocationReceipt,
    LearningOutboxRecord,
    LearningTerminalEvidence,
)
from spaces.learning.services.db.repository import PersistenceConflict
from spaces.learning.services.ingestion.chunker import chunk_blocks
from spaces.learning.services.ingestion.models import (
    LearningSourceChunk,
    LearningSourceRevision,
)
from spaces.learning.services.ingestion.normalizer import normalize_blocks
from spaces.learning.services.ingestion.parsers import (
    ParserError,
    ParsedBlock,
    parse_document,
)
from spaces.learning.services.ingestion.pipeline import IngestionPipeline
from spaces.learning.services.ingestion.repository import SourceRepository


@pytest.fixture()
def ingestion_store(tmp_path: Path):
    engine = create_engine(f"sqlite:///{tmp_path / 'ingestion.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    artifact_root = tmp_path / "artifacts"
    artifact_root.mkdir()
    repository = SourceRepository(factory, artifact_root=artifact_root)
    return factory, artifact_root, repository


def _write_pdf(path: Path, text: str) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)


def test_structured_parsers_preserve_format_specific_locators(tmp_path: Path) -> None:
    pdf_path = tmp_path / "policy.pdf"
    _write_pdf(pdf_path, "Approved AI use")

    docx_path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Governance", level=1)
    document.add_paragraph("Human review is required.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "Owner"
    document.save(docx_path)

    markdown_path = tmp_path / "policy.md"
    markdown_path.write_text(
        "# Governance\n\nHuman review is required.\n\n```python\nverify()\n```\n",
        encoding="utf-8",
    )
    text_path = tmp_path / "policy.txt"
    text_path.write_text("First paragraph.\n\nSecond paragraph.\n", encoding="utf-8")

    pdf = parse_document(pdf_path, "application/pdf")
    docx = parse_document(
        docx_path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    markdown = parse_document(markdown_path, "text/markdown")
    plain = parse_document(text_path, "text/plain")

    assert pdf.blocks[0].text == "Approved AI use"
    assert pdf.blocks[0].locator == {"page": 1}
    assert docx.blocks[1].locator == {"paragraph": 2, "heading": "Governance"}
    assert [block.text for block in docx.blocks[-2:]] == ["Control", "Owner"]
    assert docx.blocks[-1].locator == {
        "table": 1,
        "row": 1,
        "cell": 2,
        "paragraph": 1,
        "heading": "Governance",
    }
    assert markdown.blocks[1].locator == {
        "line_start": 3,
        "line_end": 3,
        "heading": "Governance",
    }
    assert any(block.metadata.get("kind") == "code" for block in markdown.blocks)
    assert [block.locator["line_start"] for block in plain.blocks] == [1, 3]


def test_normalization_and_chunk_boundaries_are_deterministic() -> None:
    blocks = [
        ParsedBlock(
            text="  Human\t review   is required.  " * 5,
            locator={"page": 4, "heading": "Controls"},
            metadata={"language": "en"},
        )
    ]

    normalized = normalize_blocks(blocks)
    first = chunk_blocks(normalized, maximum_characters=72)
    replay = chunk_blocks(normalized, maximum_characters=72)

    assert first == replay
    assert len(first) > 1
    assert all(len(chunk.content) <= 72 for chunk in first)
    assert all(
        chunk.content_hash == hashlib.sha256(chunk.content.encode()).hexdigest()
        for chunk in first
    )
    assert [chunk.locator["chunk_start"] for chunk in first] == sorted(
        chunk.locator["chunk_start"] for chunk in first
    )
    for chunk in first:
        start = int(chunk.locator["chunk_start"])
        end = int(chunk.locator["chunk_end"])
        assert normalized[0].text[start:end] == chunk.content
        assert chunk.locator["coordinate_space"] == "normalized-block-v2"
        assert chunk.locator["normalization_version"] == "2"


def test_code_normalization_preserves_unicode_and_indentation() -> None:
    block = ParsedBlock(
        text="    label = '①'  \r\n        return label\t\r\n",
        locator={"line_start": 1, "line_end": 2},
        metadata={"kind": "code"},
    )

    normalized = normalize_blocks([block])

    assert normalized[0].text == "    label = '①'\n        return label"


def test_parser_rejects_documents_over_resource_limits(tmp_path: Path) -> None:
    path = tmp_path / "large.txt"
    path.write_text("x" * 65, encoding="utf-8")

    with pytest.raises(ParserError, match="parser_resource_limit"):
        parse_document(path, "text/plain", maximum_document_characters=64)


def test_parser_timeout_terminates_the_isolated_worker(tmp_path: Path) -> None:
    path = tmp_path / "large.md"
    path.write_text("# Heading\n\n" + ("content " * 150_000), encoding="utf-8")
    existing_children = {child.pid for child in active_children()}
    started = time.monotonic()

    with pytest.raises(ParserError, match="parser_resource_limit"):
        parse_document(
            path,
            "text/markdown",
            maximum_document_characters=2_000_000,
            maximum_parse_seconds=0.01,
        )

    assert time.monotonic() - started < 5
    assert {
        child.pid for child in active_children() if child.pid not in existing_children
    } == set()


def test_pipeline_copies_original_and_replays_duplicate_without_new_rows(
    tmp_path: Path, ingestion_store
) -> None:
    session_factory, artifact_root, repository = ingestion_store
    source_path = tmp_path / "policy.txt"
    source_path.write_text("Approved use.\n\nHuman review required.", encoding="utf-8")
    pipeline = IngestionPipeline(repository, artifact_root=artifact_root)
    values = {
        "source_path": source_path,
        "source_id": str(uuid4()),
        "course_id": str(uuid4()),
        "title": "AI policy",
        "expected_revision": 0,
        "idempotency_key": "ingest-policy-r1",
        "media_type": "text/plain",
    }

    first = pipeline.ingest(**values)
    replay = pipeline.ingest(**values)
    with pytest.raises(PersistenceConflict, match="idempotency"):
        pipeline.ingest(**{**values, "expected_revision": 1})
    duplicate = pipeline.ingest(
        **{**values, "idempotency_key": "ingest-policy-duplicate"}
    )

    assert first.state == "stored"
    assert replay == first
    assert duplicate.state == "duplicate"
    assert duplicate.record == first.record
    source_path.write_text("Different content.", encoding="utf-8")
    with pytest.raises(PersistenceConflict, match="idempotency"):
        pipeline.ingest(
            **{**values, "idempotency_key": "ingest-policy-duplicate"}
        )
    source_path.write_text("Approved use.\n\nHuman review required.", encoding="utf-8")
    with pytest.raises(PersistenceConflict, match="identity"):
        pipeline.ingest(
            **{
                **values,
                "title": "Changed identity",
                "idempotency_key": "ingest-policy-wrong-identity",
            }
        )
    assert first.record is not None
    assert (artifact_root / first.relative_path).is_file()
    with session_factory() as session:
        assert session.scalar(select(func.count()).select_from(LearningSourceRevision)) == 1
        assert session.scalar(select(func.count()).select_from(LearningOutboxRecord)) == 1
        assert (
            session.scalar(select(func.count()).select_from(LearningInvocationReceipt))
            == 1
        )


def test_parser_failure_persists_failed_revision_without_chunks_or_outbox(
    tmp_path: Path, ingestion_store
) -> None:
    session_factory, artifact_root, repository = ingestion_store
    source_path = tmp_path / "broken.pdf"
    source_path.write_bytes(b"not a pdf")
    pipeline = IngestionPipeline(repository, artifact_root=artifact_root)
    source_id = str(uuid4())
    course_id = str(uuid4())

    result = pipeline.ingest(
        source_path=source_path,
        source_id=source_id,
        course_id=course_id,
        title="Broken source",
        expected_revision=0,
        idempotency_key="broken-source-r1",
        media_type="application/pdf",
    )
    replay = pipeline.ingest(
        source_path=source_path,
        source_id=source_id,
        course_id=course_id,
        title="Broken source",
        expected_revision=0,
        idempotency_key="broken-source-r1",
        media_type="application/pdf",
    )
    with pytest.raises(PersistenceConflict, match="idempotency"):
        pipeline.ingest(
            source_path=source_path,
            source_id=source_id,
            course_id=course_id,
            title="Broken source",
            expected_revision=1,
            idempotency_key="broken-source-r1",
            media_type="application/pdf",
        )
    duplicate = pipeline.ingest(
        source_path=source_path,
        source_id=source_id,
        course_id=course_id,
        title="Broken source",
        expected_revision=0,
        idempotency_key="broken-source-duplicate",
        media_type="application/pdf",
    )

    assert result.state == "failed"
    assert replay == result
    assert duplicate == result
    assert result.error_code == "parser_failed"
    assert (artifact_root / result.relative_path).is_file()
    with session_factory() as session:
        revision = session.scalar(select(LearningSourceRevision))
        chunk_count = session.scalar(select(func.count()).select_from(LearningSourceChunk))
        outbox_count = session.scalar(select(func.count()).select_from(LearningOutboxRecord))
        evidence_count = session.scalar(
            select(func.count()).select_from(LearningTerminalEvidence)
        )
        receipt_count = session.scalar(
            select(func.count()).select_from(LearningInvocationReceipt)
        )

    assert revision is not None
    assert revision.status == "failed"
    assert chunk_count == 0
    assert outbox_count == 0
    assert evidence_count == 1
    assert receipt_count == 1


def test_pipeline_rejects_artifact_reparse_point_escape(
    tmp_path: Path, ingestion_store
) -> None:
    _, artifact_root, repository = ingestion_store
    source_id = str(uuid4())
    course_id = str(uuid4())
    redirected = tmp_path / "redirected"
    redirected.mkdir()
    course_path = artifact_root / "sources" / course_id
    course_path.mkdir(parents=True)
    link_path = course_path / source_id
    try:
        os.symlink(redirected, link_path, target_is_directory=True)
    except OSError as error:
        pytest.skip(f"directory symlink unavailable: {error}")
    source_path = tmp_path / "policy.txt"
    source_path.write_text("Approved use.", encoding="utf-8")
    pipeline = IngestionPipeline(repository, artifact_root=artifact_root)

    with pytest.raises(ValueError, match="reparse|artifact root"):
        pipeline.ingest(
            source_path=source_path,
            source_id=source_id,
            course_id=course_id,
            title="AI policy",
            expected_revision=0,
            idempotency_key="reparse-escape",
            media_type="text/plain",
        )

    assert list(redirected.iterdir()) == []
