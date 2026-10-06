"""Text aus hochgeladenen Unterlagen (PDF, DOCX, TXT, MD) fuer den Gestaltungs-Agenten.

Rein: kein Modell, kein Netz, keine Dateien. `text_aus` wirft nie; was nicht
lesbar ist (kaputt, verschluesselt, leer, fremde Endung), ergibt "".
Grosse Eingaben werden schon beim Lesen begrenzt, nicht erst danach.
"""
import io
import os
import zipfile

MAX_ZEICHEN = 20_000
# Beim Lesen wird weit ueber MAX_ZEICHEN hinaus nichts mehr gesammelt.
LESE_GRENZE = 4 * MAX_ZEICHEN
_MAX_SEITEN = 200
_VERMERK = " [gekürzt]"
_MIN_ANTEIL = 200
# Zip-Bomben-Schutz: python-docx parst word/document.xml komplett (~11 MB RAM je MB XML).
_MAX_DOCUMENT_XML = 20 * 1024 * 1024
_MAX_ENTPACKT = 60 * 1024 * 1024


def _normal(text: str) -> str:
    return " ".join(text.split())


def _pdf_text(roh: bytes) -> str:
    from pypdf import PdfReader

    leser = PdfReader(io.BytesIO(roh))
    if leser.is_encrypted and not leser.decrypt(""):
        return ""
    teile: list[str] = []
    laenge = 0
    for seite in leser.pages[:_MAX_SEITEN]:
        t = (seite.extract_text() or "")[: max(0, LESE_GRENZE - laenge)]
        teile.append(t)
        laenge += len(t)
        if laenge >= LESE_GRENZE:
            break
    return " ".join(teile)


def _docx_text(roh: bytes) -> str:
    from docx import Document

    with zipfile.ZipFile(io.BytesIO(roh)) as z:
        infos = z.infolist()
        if sum(i.file_size for i in infos) > _MAX_ENTPACKT:
            return ""
        if any(i.filename == "word/document.xml" and i.file_size > _MAX_DOCUMENT_XML for i in infos):
            return ""
    dok = Document(io.BytesIO(roh))
    teile: list[str] = []
    laenge = 0

    def sammeln(t: str) -> bool:
        nonlocal laenge
        teile.append(t)
        laenge += len(t)
        return laenge >= LESE_GRENZE

    for absatz in dok.paragraphs:
        if sammeln(absatz.text):
            return " ".join(teile)
    for tabelle in dok.tables:
        for zeile in tabelle.rows:
            for zelle in zeile.cells:
                if sammeln(zelle.text):
                    return " ".join(teile)
    return " ".join(teile)


def _klartext(roh: bytes) -> str:
    roh = roh[: LESE_GRENZE * 4]
    if roh.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return roh.decode("utf-16")
        except UnicodeDecodeError:
            pass
    try:
        return roh.decode("utf-8-sig")
    except UnicodeDecodeError:
        return roh.decode("latin-1")


_LESER = {
    ".pdf": _pdf_text,
    ".docx": _docx_text,
    ".txt": _klartext,
    ".md": _klartext,
}


def text_aus(name: str, roh: bytes) -> str:
    """Lesbarer Text einer Unterlage, Whitespace normalisiert; "" wenn keiner. Wirft nie."""
    try:
        leser = _LESER.get(os.path.splitext(str(name))[1].lower())
        if leser is None:
            return ""
        return _normal(leser(roh))[:LESE_GRENZE]
    except Exception:  # Parser-Fehler jeder Art (pypdf, zipfile, XML) = nicht lesbar
        return ""


def unterlagen_text(
    dateien: list[tuple[str, bytes]], max_zeichen: int = MAX_ZEICHEN
) -> tuple[str, list[str]]:
    """("Unterlage: a.pdf\\n…\\n\\nUnterlage: b.docx\\n…", hinweise).

    Ueberlange Summen werden anteilig gekuerzt, jede gekuerzte Unterlage endet auf
    "[gekürzt]". Dateien ohne Text erscheinen nur als Hinweis.
    """
    gelesen: list[tuple[str, str]] = []
    hinweise: list[str] = []
    for name, roh in dateien:
        t = text_aus(name, roh)
        if t:
            gelesen.append((name, t))
        else:
            hinweise.append(f"{name} hat keinen lesbaren Text")
    koepfe = [f"Unterlage: {name}\n" for name, _ in gelesen]
    frei = max_zeichen - sum(len(k) for k in koepfe) - 2 * max(0, len(gelesen) - 1)
    gesamt = sum(len(t) for _, t in gelesen)
    anteile = [len(t) for _, t in gelesen]
    if gesamt > frei:
        # Vermerk fuer jede Unterlage einplanen, Mindestanteil vorweg, Rest anteilig.
        budget = max(0, frei - len(_VERMERK) * len(gelesen))
        mind = min(_MIN_ANTEIL, budget // len(gelesen))
        basis = [min(n, mind) for n in anteile]
        rest = budget - sum(basis)
        offen = sum(n - b for n, b in zip(anteile, basis))
        anteile = [b + rest * (n - b) // offen if offen else b for n, b in zip(anteile, basis)]
    abschnitte: list[str] = []
    for (name, t), kopf, anteil in zip(gelesen, koepfe, anteile):
        if anteil < len(t):
            t = t[:anteil].rstrip() + _VERMERK
        abschnitte.append(kopf + t)
    return "\n\n".join(abschnitte), hinweise