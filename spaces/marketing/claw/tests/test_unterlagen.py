"""unterlagen: Text aus PDF/DOCX/TXT/MD fuer den Gestaltungs-Agenten.

text_aus wirft nie; Unlesbares (kaputt, verschluesselt, leer) ergibt "" und
unterlagen_text macht daraus einen Hinweis.
"""
import io
import unittest
import zipfile
from unittest import mock

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from spaces.marketing.claw import unterlagen
from spaces.marketing.claw.unterlagen import MAX_ZEICHEN, text_aus, unterlagen_text


def _pdf(seiten):
    puffer = io.BytesIO()
    c = canvas.Canvas(puffer)
    for s in seiten:
        c.drawString(72, 700, s)
        c.showPage()
    c.save()
    return puffer.getvalue()


def _pdf_verschluesselt(passwort):
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(_pdf(["Geheim 123"]))))
    w.encrypt(passwort)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def _docx(absaetze=(), zellen=()):
    d = Document()
    for a in absaetze:
        d.add_paragraph(a)
    if zellen:
        t = d.add_table(rows=1, cols=len(zellen))
        for i, z in enumerate(zellen):
            t.rows[0].cells[i].text = z
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


class TextAus(unittest.TestCase):
    def test_pdf(self):
        self.assertIn("Sommerangebot 9,90", text_aus("a.pdf", _pdf(["Sommerangebot 9,90"])))

    def test_pdf_mehrere_seiten(self):
        t = text_aus("a.PDF", _pdf(["Seite eins", "Seite zwei"]))
        self.assertIn("Seite eins", t)
        self.assertIn("Seite zwei", t)

    def test_docx_absaetze_und_tabelle(self):
        t = text_aus("a.docx", _docx(["Hallo Welt"], ["Preis", "4,50"]))
        self.assertIn("Hallo Welt", t)
        self.assertIn("Preis", t)
        self.assertIn("4,50", t)

    def test_txt_utf8_und_latin1(self):
        self.assertEqual(text_aus("a.txt", "Käse  und\n\n Brot".encode("utf-8")), "Käse und Brot")
        self.assertEqual(text_aus("a.txt", "Käse".encode("latin-1")), "Käse")

    def test_md(self):
        self.assertEqual(text_aus("a.md", b"# Titel\n\ntext"), "# Titel text")

    def test_kaputtes_pdf(self):
        self.assertEqual(text_aus("a.pdf", b"%PDF-1.4 das ist kein pdf"), "")
        self.assertEqual(text_aus("a.pdf", b""), "")

    def test_verschluesseltes_pdf(self):
        self.assertEqual(text_aus("a.pdf", _pdf_verschluesselt("streng")), "")

    def test_pdf_mit_leerem_passwort_wird_gelesen(self):
        self.assertIn("Geheim 123", text_aus("a.pdf", _pdf_verschluesselt("")))

    def test_leeres_docx(self):
        self.assertEqual(text_aus("a.docx", _docx()), "")

    def test_kaputtes_docx(self):
        self.assertEqual(text_aus("a.docx", b"PK\x03\x04 kaputt"), "")
        self.assertEqual(text_aus("a.docx", b"kein zip"), "")

    def test_utf8_bom_und_utf16(self):
        self.assertEqual(text_aus("a.txt", "\ufeffKäse".encode("utf-8")), "Käse")
        self.assertEqual(text_aus("a.txt", "Käse ß".encode("utf-16")), "Käse ß")
        self.assertEqual(text_aus("a.md", "Käse".encode("utf-16-be").join([b"\xfe\xff", b""])), "Käse")

    def test_docx_zip_bombe(self):
        puffer = io.BytesIO()
        with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", "<a>" + "0" * 1_000_000 + "</a>")
        bombe = puffer.getvalue()
        self.assertLess(len(bombe), 20_000)
        with mock.patch.object(unterlagen, "_MAX_DOCUMENT_XML", 100_000):
            self.assertEqual(text_aus("a.docx", bombe), "")
        # Gesamtgroesse aller Mitglieder
        ok = _docx(["Hallo"])
        with mock.patch.object(unterlagen, "_MAX_ENTPACKT", 100):
            self.assertEqual(text_aus("a.docx", ok), "")
        self.assertIn("Hallo", text_aus("a.docx", ok))

    def test_pdf_riesige_einzelseite(self):
        with mock.patch.object(unterlagen, "LESE_GRENZE", 50):
            t = text_aus("a.pdf", _pdf(["w" * 400]))
        self.assertLessEqual(len(t), 50)

    def test_unbekannte_endung(self):
        self.assertEqual(text_aus("a.exe", b"text"), "")
        self.assertEqual(text_aus("ohne", b"text"), "")

    def test_riesige_eingabe_wird_begrenzt(self):
        t = text_aus("a.txt", b"wort " * 2_000_000)
        self.assertLessEqual(len(t), unterlagen.LESE_GRENZE)

    def test_viele_pdf_seiten_stoppen_frueh(self):
        t = text_aus("a.pdf", _pdf(["x" * 3000] * 40))
        self.assertGreater(len(t), 0)
        self.assertLessEqual(len(t), unterlagen.LESE_GRENZE)

    def test_viele_docx_absaetze_begrenzt(self):
        t = text_aus("a.docx", _docx(["y" * 1000] * 500))
        self.assertLessEqual(len(t), unterlagen.LESE_GRENZE + 1000)


class UnterlagenText(unittest.TestCase):
    def test_zwei_dateien_mit_kopf(self):
        text, hinweise = unterlagen_text([("a.txt", b"eins"), ("b.md", b"zwei")])
        self.assertEqual(text, "Unterlage: a.txt\neins\n\nUnterlage: b.md\nzwei")
        self.assertEqual(hinweise, [])

    def test_datei_ohne_text_gibt_hinweis(self):
        text, hinweise = unterlagen_text([
            ("kaputt.pdf", b"xx"), ("leer.docx", _docx()), ("ok.txt", b"gut"),
            ("zu.pdf", _pdf_verschluesselt("pw")),
        ])
        self.assertEqual(text, "Unterlage: ok.txt\ngut")
        self.assertEqual(hinweise, [
            "kaputt.pdf hat keinen lesbaren Text",
            "leer.docx hat keinen lesbaren Text",
            "zu.pdf hat keinen lesbaren Text",
        ])

    def test_leere_liste(self):
        self.assertEqual(unterlagen_text([]), ("", []))

    def test_kuerzung_anteilig_bei_drei_dateien(self):
        dateien = [("a.txt", b"a" * 30_000), ("b.txt", b"b" * 15_000), ("c.txt", b"c" * 5_000)]
        text, hinweise = unterlagen_text(dateien)
        self.assertEqual(text.count("[gekürzt]"), 3)
        teile = text.split("\n\nUnterlage: ")
        la, lb, lc = (t.count(x) for t, x in zip(teile, "abc"))
        self.assertGreater(la, lb)
        self.assertGreater(lb, lc)
        self.assertGreater(lc, 1500)
        self.assertLessEqual(len(text), MAX_ZEICHEN)
        self.assertEqual(hinweise, [])

    def test_viele_kleine_dateien_gesamt_inklusive_koepfe(self):
        dateien = [(f"datei{i}.txt", b"t" * 900) for i in range(40)]
        text, _ = unterlagen_text(dateien)
        self.assertLessEqual(len(text), MAX_ZEICHEN)
        self.assertEqual(text.count("Unterlage: "), 40)

    def test_kleine_datei_neben_riesiger_behaelt_text(self):
        text, _ = unterlagen_text([("gross.txt", b"g" * 100_000), ("klein.txt", b"Preis 4,50")])
        self.assertLessEqual(len(text), MAX_ZEICHEN)
        self.assertIn("Unterlage: klein.txt\nPreis 4,50", text)
        self.assertEqual(text.count("[gekürzt]"), 1)

    def test_kein_vermerk_wenn_es_passt(self):
        text, _ = unterlagen_text([("a.txt", b"a" * 100)])
        self.assertNotIn("[gekürzt]", text)

    def test_eigene_grenze(self):
        text, _ = unterlagen_text([("a.txt", b"a" * 500)], max_zeichen=100)
        self.assertIn("[gekürzt]", text)
        self.assertLessEqual(text.split("\n", 1)[1].count("a"), 100)


if __name__ == "__main__":
    unittest.main()

