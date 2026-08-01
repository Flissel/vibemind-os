"""anschreiben-generator: Generiert ein Nachreichungs-Anschreiben als docx.

Liest Status-Report (aus status-abgleich) + stammdaten.yaml, baut ein
formales Behörden-Anschreiben programmatisch mit python-docx. Listet die
offenen Forderungen und erklärt N/A-Fälle.

Aufruf:
    python skills/buergergeld/anschreiben-generator/run.py <status_report.yaml>
    python skills/buergergeld/anschreiben-generator/run.py <status_report.yaml> --typ nachreichung
    python skills/buergergeld/anschreiben-generator/run.py <status_report.yaml> --dry-run
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml

_SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_SKILL_ROOT))


def _load_lib_module(name: str, file: str):
    spec = importlib.util.spec_from_file_location(
        f"buergergeld_{name}", _SKILL_ROOT / "_lib" / file
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


_timeline = _load_lib_module("timeline_helper", "timeline_helper.py")

append_event = _timeline.append_event
buergergeld_root = _timeline.buergergeld_root
get_stammdaten = _timeline.get_stammdaten


MONATE = {
    1: "Januar", 2: "Februar", 3: "März", 4: "April", 5: "Mai", 6: "Juni",
    7: "Juli", 8: "August", 9: "September", 10: "Oktober", 11: "November", 12: "Dezember",
}


def fmt_datum_lang(d: date | None = None) -> str:
    d = d or date.today()
    if isinstance(d, str):
        try:
            d = datetime.strptime(d, "%Y-%m-%d").date()
        except ValueError:
            return d
    return f"{d.day}. {MONATE[d.month]} {d.year}"


def load_report(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def build_docx(report: dict[str, Any], stammdaten: dict[str, Any], typ: str) -> Any:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    person = stammdaten.get("person", {})
    adresse = stammdaten.get("adresse", {})
    jobcenter = stammdaten.get("jobcenter", {})

    name = f"{person.get('vorname', '')} {person.get('nachname', '')}".strip()
    bg_nr = jobcenter.get("bg_nummer", "")
    kunden_nr = jobcenter.get("kundennummer", "")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def para(text="", *, bold=False, align=None, size=None, space_after=6):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    # === Absender ===
    para(name, bold=True, space_after=0)
    para(adresse.get("strasse", ""), space_after=0)
    para(f"{adresse.get('plz', '')} {adresse.get('ort', '')}".strip(), space_after=0)
    if person.get("geburtsdatum"):
        gd = person["geburtsdatum"]
        gd_str = fmt_datum_lang(gd) if not isinstance(gd, str) else gd
        para(f"Geburtsdatum: {gd_str}", space_after=0)
    if bg_nr:
        para(f"BG-Nummer: {bg_nr}", space_after=0)
    if kunden_nr:
        para(f"Kundennummer: {kunden_nr}", space_after=12)
    else:
        para(space_after=12)

    # === Empfänger ===
    para(jobcenter.get("name", "Jobcenter"), space_after=0)
    jc_adr = jobcenter.get("adresse", {})
    if isinstance(jc_adr, dict):
        para(jc_adr.get("strasse", ""), space_after=0)
        para(f"{jc_adr.get('plz', '')} {jc_adr.get('ort', '')}".strip(), space_after=12)
    else:
        para(str(jc_adr), space_after=12)

    # === Ort, Datum ===
    para(f"{adresse.get('ort', 'München')}, den {fmt_datum_lang()}", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

    # === Betreff ===
    bezug_betreff = report.get("betreff", "")
    datum_schreiben = report.get("datum_schreiben")
    datum_str = fmt_datum_lang(datum_schreiben) if datum_schreiben else ""
    betreff = "Nachreichung von Unterlagen zum Bürgergeld-Antrag"
    if datum_str:
        betreff += f" — Ihr Schreiben vom {datum_str}"
    para(f"Betreff: {betreff}", bold=True, space_after=12)

    # === Anrede ===
    para("Sehr geehrte Damen und Herren,", space_after=6)

    # === Einleitung ===
    offen = [f for f in report["forderungen"] if f["klassifikation"] == "offen"]
    erfuellt = [f for f in report["forderungen"] if f["klassifikation"] == "erfuellt"]
    na_artige = [f for f in report["forderungen"]
                 if f["klassifikation"] in ("offen_optional",)]

    intro = (
        f"in Bezug auf Ihr Schreiben{' vom ' + datum_str if datum_str else ''} reiche ich "
        f"hiermit die noch ausstehenden Unterlagen nach beziehungsweise nehme zu den "
        f"angeforderten Punkten Stellung."
    )
    para(intro, space_after=12)

    # === 1. Beigefügte / nachgereichte Unterlagen ===
    if offen:
        para("1. Nachgereichte Unterlagen", bold=True, space_after=6)
        para("Folgende Unterlagen reiche ich mit diesem Schreiben nach:", space_after=6)
        for f in offen:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f["forderung_titel"])
            p.paragraph_format.space_after = Pt(3)
        para(space_after=6)

    # === 2. Erläuterungen zu nicht zutreffenden Punkten ===
    if na_artige:
        para("2. Erläuterungen zu weiteren Punkten", bold=True, space_after=6)
        para(
            "Zu den folgenden, in Ihrem Schreiben (teils mit dem Zusatz „soweit "
            "vorhanden\" oder „gegebenenfalls\") genannten Punkten gebe ich an:",
            space_after=6,
        )
        for f in na_artige:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{f['forderung_titel']}: ").bold = True
            p.add_run("liegt nicht vor / nicht zutreffend.")
            p.paragraph_format.space_after = Pt(3)
        para(space_after=6)

    # === 3. Bereits eingereichte Unterlagen (Hinweis) ===
    if erfuellt:
        para("3. Bereits eingereichte Unterlagen", bold=True, space_after=6)
        para(
            f"Die übrigen in Ihrem Schreiben angeforderten Unterlagen "
            f"({len(erfuellt)} Positionen) habe ich bereits mit meiner vorherigen "
            f"Übermittlung eingereicht. Eine vollständige Aufstellung liegt diesem "
            f"Schreiben als Anlage bei.",
            space_after=12,
        )

    # === Schluss ===
    para(
        "Für Rückfragen oder die Anforderung weiterer Nachweise stehe ich Ihnen "
        "selbstverständlich jederzeit zur Verfügung.",
        space_after=12,
    )
    para("Mit freundlichen Grüßen", space_after=24)
    para("_______________________________", space_after=0)
    para(name, space_after=0)

    return doc


def run(report_path: Path, typ: str = "nachreichung", dry: bool = False) -> dict[str, Any]:
    if not report_path.is_absolute():
        candidate = buergergeld_root() / report_path
        report_path = candidate if candidate.exists() else Path.cwd() / report_path

    if not report_path.exists():
        raise FileNotFoundError(f"Status-Report nicht gefunden: {report_path}")

    report = load_report(report_path)
    stammdaten = get_stammdaten()

    offen = [f for f in report["forderungen"] if f["klassifikation"] == "offen"]

    result = {
        "report": str(report_path),
        "typ": typ,
        "n_offen": len(offen),
        "offene_titel": [f["forderung_titel"] for f in offen],
    }

    if dry:
        result["dry_run"] = True
        return result

    doc = build_docx(report, stammdaten, typ)

    today = date.today().isoformat()
    ausgang = buergergeld_root() / "ausgang"
    ausgang.mkdir(exist_ok=True)
    out_path = ausgang / f"{today}_Anschreiben_{typ.capitalize()}.docx"
    doc.save(str(out_path))

    result["output"] = str(out_path)

    append_event(
        typ="anschreiben_generiert",
        titel=f"Anschreiben generiert ({typ}): {len(offen)} nachzureichende Positionen",
        quelle="skill:buergergeld-anschreiben-generator",
        anschreiben_typ=typ,
        n_offen=len(offen),
        output=str(out_path.relative_to(buergergeld_root())).replace("\\", "/")
            if str(out_path).startswith(str(buergergeld_root())) else str(out_path),
    )

    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("report", help="Pfad zum status_report.yaml")
    parser.add_argument("--typ", default="nachreichung",
                        choices=["nachreichung", "mietschuldenuebernahme", "veraenderungsanzeige"])
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    result = run(Path(args.report), typ=args.typ, dry=args.dry_run)
    print(json.dumps(result, indent=2, default=str, ensure_ascii=False))
